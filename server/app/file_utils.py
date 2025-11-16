import io
import pdfplumber
from docx import Document


def extract_pdf_text(pdf_bytes: bytes) -> str:
    parts = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            parts.append(text or "")
    return "\n\n".join(parts)


def extract_docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = []

    # собираем текст параграфов
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())

    # если есть таблицы — тоже добавим
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            parts.append(row_text)

    return "\n\n".join(parts)


def extract_text(file_bytes: bytes, ext: str) -> str:
    ext = ext.lower()

    if ext.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    if ext.endswith(".docx"):
        return extract_docx_text(file_bytes)

    raise ValueError("Unsupported file")
